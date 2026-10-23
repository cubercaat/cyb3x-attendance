"""Hisobotlar — chiroyli Excel (.xlsx) va jonli HTML jadval."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app import models as m

_GREEN = PatternFill("solid", fgColor="D6F5E1")
_YELLOW = PatternFill("solid", fgColor="FFF4D6")
_RED = PatternFill("solid", fgColor="FBDADA")
_HEAD = PatternFill("solid", fgColor="0E7C66")
_TITLE = PatternFill("solid", fgColor="0B3D2E")
_thin = Side(style="thin", color="CCCCCC")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

STATUS_UZ = {"present": "✔ Keldi", "late": "⏰ Kech", "absent": "✘ Kelmadi"}


def build_xlsx(summary: dict, group_name: str) -> bytes:
    s: m.AttSession = summary["session"]
    wb = Workbook()
    ws = wb.active
    ws.title = f"{s.para}-para"

    # Sarlavha
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = f"CYB3X DAVOMAT  —  {group_name}  —  {s.para}-para  ({s.opened_at.strftime('%Y-%m-%d')})"
    c.font = Font(bold=True, color="FFFFFF", size=13)
    c.fill = _TITLE
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    # Jami qator
    ws.merge_cells("A2:E2")
    t = ws["A2"]
    total = summary["total"] or 1
    pct = round(summary["present"] * 100 / total)
    t.value = (f"Jami: {summary['total']}   |   ✔ Keldi: {summary['present']}   "
               f"|   ⏰ Kech: {summary['late']}   |   ✘ Kelmadi: {summary['absent']}   "
               f"|   Davomat: {pct}%")
    t.font = Font(bold=True, size=11)
    t.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 20

    headers = ["№", "F.I.Sh.", "Holati", "Usul", "Vaqt"]
    for j, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=j, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    for i, row in enumerate(summary["rows"], start=1):
        r = i + 3
        vals = [i, row["name"], STATUS_UZ.get(row["status"], row["status"]), row["method"], row["time"]]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="center" if j != 2 else "left", vertical="center")
        fill = {"present": _GREEN, "late": _YELLOW, "absent": _RED}.get(row["status"])
        if fill:
            for j in range(1, 6):
                ws.cell(row=r, column=j).fill = fill

    widths = [5, 32, 14, 10, 12]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
